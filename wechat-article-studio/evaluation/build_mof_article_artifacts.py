from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "evaluation" / "mof-solid-electrolyte-investor-article.md"
HTML_OUTPUT = ROOT / "evaluation" / "mof-solid-electrolyte-investor-article.html"
DOCX_OUTPUT = ROOT / "evaluation" / "mof-solid-electrolyte-investor-article.docx"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
BODY = RGBColor(48, 48, 48)
GRAY = RGBColor(85, 85, 85)
MUTED = RGBColor(115, 115, 115)

READER_FORBIDDEN_FIELDS = (
    "human_curated", "realtime_research", "sourceOrigin", "claim:", "source:",
    "review_required", "candidate", "Bridge", "HTTP 504", "research_timeout",
    "EvidencePacket", "ArgumentMap", "WriterRequest", "WechatPackage", "clientRunId", "allReady",
)


def parse_blocks(markdown: str):
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks = []
    paragraph = []
    quote = []
    list_items = []
    list_kind = None

    def flush_paragraph():
        nonlocal paragraph
        if paragraph:
            blocks.append(("p", " ".join(part.strip() for part in paragraph).strip()))
            paragraph = []

    def flush_quote():
        nonlocal quote
        if quote:
            blocks.append(("quote", " ".join(part.strip() for part in quote).strip()))
            quote = []

    def flush_list():
        nonlocal list_items, list_kind
        if list_items:
            blocks.append((list_kind, list_items))
            list_items = []
            list_kind = None

    for raw in lines:
        line = raw.strip()
        if not line:
            flush_paragraph()
            flush_quote()
            flush_list()
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph(); flush_quote(); flush_list()
            blocks.append((f"h{len(heading.group(1))}", heading.group(2).strip()))
            continue
        if line.startswith(">"):
            flush_paragraph(); flush_list()
            quote.append(line[1:].strip())
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph(); flush_quote()
            kind = "ul" if bullet else "ol"
            if list_kind and list_kind != kind:
                flush_list()
            list_kind = kind
            list_items.append((bullet or numbered).group(1).strip())
            continue
        flush_quote(); flush_list()
        paragraph.append(line)
    flush_paragraph(); flush_quote(); flush_list()
    return blocks


def assert_reader_markdown(markdown: str):
    leaked = [field for field in READER_FORBIDDEN_FIELDS if field in markdown]
    if leaked:
        raise ValueError(f"Reader Markdown contains internal fields: {', '.join(leaked)}")
    if re.search(r"\[S\d+\]", markdown):
        raise ValueError("Reader Markdown must use numeric reference markers such as [1], not [S1]")
    if "## 来源与来源类型" in markdown:
        raise ValueError("Reader Markdown must use the heading 参考文献")


def inline_html(text: str) -> str:
    escaped = html.escape(text, quote=False)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)
    escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
    escaped = re.sub(
        r"(https?://[^\s<]+)",
        lambda match: f'<a href="{match.group(1)}">{match.group(1)}</a>',
        escaped,
    )
    return escaped


def render_html(blocks) -> str:
    body = []
    for kind, value in blocks:
        if kind.startswith("h"):
            level = kind[1:]
            body.append(f"<{kind}>{inline_html(value)}</{kind}>")
        elif kind == "p":
            body.append(f"<p>{inline_html(value)}</p>")
        elif kind == "quote":
            body.append(f"<blockquote>{inline_html(value)}</blockquote>")
        elif kind in {"ul", "ol"}:
            body.append(f"<{kind}>" + "".join(f"<li>{inline_html(item)}</li>" for item in value) + f"</{kind}>")
    return """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>MOF 用于固态电解质，离商业化还有多远？</title>
  <style>
    :root { color-scheme: light; --ink:#243447; --muted:#657487; --blue:#2e74b5; --line:#d9e0e8; --warn:#fff8e7; --warn-border:#d5a72b; }
    * { box-sizing:border-box; }
    body { margin:0; background:#f5f7fa; color:var(--ink); font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Microsoft YaHei",Arial,sans-serif; line-height:1.85; }
    main { max-width:860px; margin:40px auto; padding:48px 64px 64px; background:#fff; box-shadow:0 8px 32px rgba(26,45,68,.08); }
    h1 { margin:0 0 20px; color:#0b2545; font-size:34px; line-height:1.3; letter-spacing:.01em; }
    h2 { margin:40px 0 14px; color:var(--blue); font-size:23px; line-height:1.4; border-left:4px solid var(--blue); padding-left:12px; }
    h3 { margin:28px 0 10px; color:#1f4d78; font-size:18px; line-height:1.45; }
    p { margin:0 0 14px; font-size:16px; }
    blockquote { margin:18px 0; padding:14px 18px; border-left:4px solid var(--warn-border); background:var(--warn); color:#5c4b22; font-size:15px; }
    ol, ul { margin:8px 0 18px 24px; padding-left:20px; }
    li { margin:6px 0; font-size:16px; }
    code { padding:1px 4px; background:#eef2f6; border-radius:3px; font-size:.9em; }
    a { color:#1264a3; word-break:break-all; }
    strong { color:#183b61; }
    h1 + blockquote { margin-top:0; }
    @media (max-width:700px) { main { margin:0; padding:28px 20px 40px; box-shadow:none; } h1 { font-size:27px; } h2 { font-size:20px; } p, li { font-size:15px; } }
  </style>
</head>
<body><main>
""" + "\n".join(body) + "\n</main></body></html>"


def set_run_font(run, name="Calibri", size=11, color=BODY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", size=11, color=BODY, bold=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths=(9360,), indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    if tbl_w.getparent() is None: tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths))); tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    if tbl_ind.getparent() is None: tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent)); tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW")) or OxmlElement("w:tcW")
            if tc_w.getparent() is None: cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index])); tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end]); set_run_font(run, size=9, color=MUTED)


def add_inline_runs(paragraph, text, *, size=11, color=BODY):
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); set_run_font(run, size=size, color=NAVY, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); set_run_font(run, size=size - .5, color=GRAY)
        else:
            run = paragraph.add_run(part); set_run_font(run, size=size, color=color)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(1); section.right_margin = Inches(1)
    section.bottom_margin = Inches(1); section.left_margin = Inches(1)
    section.header_distance = Inches(.492); section.footer_distance = Inches(.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=BODY)
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10; normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title = doc.styles["Title"]
    set_style_font(title, size=23, color=RGBColor(0, 0, 0), bold=True)
    title.paragraph_format.space_before = Pt(0); title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, size=14, color=GRAY)
    subtitle.paragraph_format.space_before = Pt(0); subtitle.paragraph_format.space_after = Pt(16)
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = doc.styles[name]; set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]; set_style_font(style, size=11, color=BODY)
        style.paragraph_format.space_after = Pt(4); style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.left_indent = Inches(.5); style.paragraph_format.first_line_indent = Inches(-.25)
    if "Research Status" not in [style.name for style in doc.styles]:
        callout = doc.styles.add_style("Research Status", 1)
        set_style_font(callout, size=10.5, color=RGBColor(92, 75, 34))
        callout.paragraph_format.left_indent = Inches(.2); callout.paragraph_format.right_indent = Inches(.2)
        callout.paragraph_format.space_before = Pt(4); callout.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("MOF 固态电解质商业化评估"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("证据优先研究 · 第 "), size=9, color=MUTED); set_page_number(footer)
    set_run_font(footer.add_run(" 页"), size=9, color=MUTED)


def add_docx_blocks(doc, blocks):
    for kind, value in blocks:
        if kind.startswith("h"):
            level = int(kind[1:])
            p = doc.add_heading(level=level); add_inline_runs(p, value, size={1:16, 2:13, 3:12}[level], color=BLUE if level < 3 else DARK_BLUE)
        elif kind == "p":
            p = doc.add_paragraph(); add_inline_runs(p, value)
        elif kind == "quote":
            p = doc.add_paragraph(style="Research Status"); add_inline_runs(p, value, size=10.5, color=RGBColor(92, 75, 34))
            p.paragraph_format.left_indent = Inches(.2)
            p.paragraph_format.right_indent = Inches(.2)
        elif kind in {"ul", "ol"}:
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in value:
                p = doc.add_paragraph(style=style); add_inline_runs(p, item)


def build_docx(blocks):
    doc = Document(); configure_document(doc)
    doc.core_properties.title = "MOF 用于固态电解质，离商业化还有多远？"
    doc.core_properties.subject = "MOF 固态电解质商业化前景分析"
    doc.core_properties.author = "公众号文章工作台"
    doc.core_properties.comments = "依据公开论文、综述、专利和企业资料整理。"
    add_docx_blocks(doc, blocks)
    doc.save(DOCX_OUTPUT)


def main():
    markdown = INPUT.read_text(encoding="utf-8")
    assert_reader_markdown(markdown)
    blocks = parse_blocks(markdown)
    HTML_OUTPUT.write_text(render_html(blocks), encoding="utf-8")
    build_docx(blocks)
    print(HTML_OUTPUT)
    print(DOCX_OUTPUT)


if __name__ == "__main__":
    main()
