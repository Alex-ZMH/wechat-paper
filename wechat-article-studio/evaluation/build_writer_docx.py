from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "evaluation" / "real-writer-runs.json"
OUTPUT = ROOT / "evaluation" / "真实writer三篇公众号文章.docx"

NAVY = RGBColor(32, 55, 72)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(85, 85, 85)
MUTED = RGBColor(115, 115, 115)
BODY = RGBColor(48, 48, 48)

INTERNAL_TOKEN_PATTERN = re.compile(
    r"(?:\[(?:claim|source|evidence|internal):[^\]\r\n]+\]|"
    r"\{\{(?:claim|source|evidence|internal):[^}\r\n]+\}\}|"
    r"<!--\s*internal:[\s\S]*?-->)",
    re.IGNORECASE,
)


def clean(value: object) -> str:
    text = str(value or "")
    text = INTERNAL_TOKEN_PATTERN.sub("", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\s+([，。；：、！？,.!?])", r"\1", text)
    return text.strip()


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", size=11, color=BODY, bold=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa=(3000, 6360), indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = doc.styles["Title"]
    set_style_font(title, size=30, color=NAVY)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, size=14, color=GRAY)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("真实 writer 评测稿 · 供人工审阅")
    set_run_font(header_run, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("EvidencePacket + Codex CLI writer · 第 ")
    set_run_font(footer_run, size=9, color=MUTED)
    set_page_number(footer)
    footer_tail = footer.add_run(" 页")
    set_run_font(footer_tail, size=9, color=MUTED)


def add_labeled_paragraph(doc, label, value, *, style="Normal", color=BODY):
    paragraph = doc.add_paragraph(style=style)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=10.5, color=NAVY, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=10.5, color=color)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_cover(doc, data):
    for _ in range(5):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(12)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("真实 WRITER 评测")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("三篇公众号文章")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("结构化初稿 · 供人工审阅")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(34)
    meta_run = meta.add_run(
        f"生成方式：{data.get('provider', 'codex-cli')} / {data.get('model', 'gpt-5.6-sol')}\n"
        "证据边界：人工确认、带官方来源 URL 的 EvidencePacket\n"
        "阅读目的：检查逻辑推进、证据准确性、重复表达、文风一致性和可读性"
    )
    set_run_font(meta_run, size=10.5, color=GRAY)

    doc.add_heading("阅读说明", level=1)
    add_labeled_paragraph(doc, "这是什么：", "三篇由真实 Codex CLI writer 生成的结构化公众号初稿。")
    add_labeled_paragraph(doc, "怎么看：", "请直接阅读每篇正文；文章末尾附有来源和自动审查结果。")
    add_labeled_paragraph(doc, "重要限制：", "本轮研究接口曾超时，因此证据包由人工确认，不代表实时检索已经稳定。")
    add_labeled_paragraph(doc, "最终判断：", "ReviewReport 通过不等于高质量定稿，是否符合你的标准仍需要你签字。")


def add_review_table(doc, run_data):
    review = run_data.get("reviewReport", {})
    checks = review.get("checks", {})
    rows = [
        ("逻辑推进", checks.get("logic", False)),
        ("证据覆盖", checks.get("evidence", False)),
        ("重复检查", checks.get("repetition", False)),
        ("文风检查", checks.get("style", False)),
        ("可读性检查", checks.get("readability", False)),
        ("自动审查结果", f"{review.get('status', 'unknown')} / {run_data.get('wechatPackage', {}).get('status', 'unknown')}"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table)
    hdr = table.rows[0].cells
    hdr[0].text = "检查项"
    hdr[1].text = "结果"
    mark_header_row(table.rows[0])
    for cell in hdr:
        set_cell_shading(cell, "F2F4F7")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10, color=NAVY, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = ("通过" if value else "未通过") if isinstance(value, bool) else str(value)
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=10, color=BODY)
        set_table_geometry(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_sources(doc, run_data):
    packet = run_data.get("evidencePacket", {})
    sources = packet.get("sources", [])
    doc.add_heading("来源与证据", level=2)
    if not sources:
        add_labeled_paragraph(doc, "来源：", "本篇没有可展示的来源记录。")
        return
    for index, source in enumerate(sources, start=1):
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(6)
        title_run = paragraph.add_run(f"来源 {index} · {clean(source.get('title'))}\n")
        set_run_font(title_run, size=10.5, color=NAVY, bold=True)
        url = clean(source.get("url"))
        if url:
            url_run = paragraph.add_run(f"{url}\n")
            set_run_font(url_run, size=9.5, color=BLUE)
        excerpt_run = paragraph.add_run(f"摘录：{clean(source.get('excerpt'))}")
        set_run_font(excerpt_run, size=10, color=GRAY, italic=True)


def add_article(doc, run_data, index, total):
    if index > 1:
        doc.add_page_break()
    draft = run_data.get("draft", {})
    brief = run_data.get("brief", {})
    doc.add_paragraph(f"文章 {index} / {total}", style="Subtitle")
    title = doc.add_heading(clean(draft.get("title")), level=1)
    title.paragraph_format.keep_with_next = True
    add_labeled_paragraph(doc, "选题：", clean(brief.get("topic")))
    add_labeled_paragraph(doc, "摘要：", clean(draft.get("digest")))
    doc.add_heading("正文", level=2)
    lead = doc.add_paragraph(style="Normal")
    lead.paragraph_format.space_after = Pt(12)
    lead_run = lead.add_run(clean(draft.get("lead")))
    set_run_font(lead_run, size=11.5, color=NAVY, italic=True)

    for section in draft.get("sections", []):
        doc.add_heading(clean(section.get("heading")), level=2)
        for paragraph in section.get("paragraphs", []):
            body = doc.add_paragraph(clean(paragraph.get("text")), style="Normal")
            body.paragraph_format.widow_control = True

    doc.add_heading("结语", level=2)
    doc.add_paragraph(clean(draft.get("closingCta")), style="Normal")
    add_sources(doc, run_data)
    doc.add_heading("自动审查结果", level=2)
    add_review_table(doc, run_data)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run("编辑提示：以上结果是机器审查和初步人工记录，不代替你的最终审阅。")
    set_run_font(note_run, size=9.5, color=MUTED, italic=True)


def main():
    data = json.loads(INPUT.read_text(encoding="utf-8"))
    runs = data.get("runs", [])
    if not runs:
        raise SystemExit("real-writer-runs.json does not contain any runs")
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "真实 writer 生成的三篇公众号文章"
    doc.core_properties.subject = "结构化初稿人工审阅稿"
    doc.core_properties.author = "公众号文章工作台"
    doc.core_properties.comments = "由真实 Codex CLI writer 生成；证据包为人工确认来源。"
    add_cover(doc, data)
    doc.add_page_break()
    for index, run_data in enumerate(runs, start=1):
        add_article(doc, run_data, index, len(runs))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
