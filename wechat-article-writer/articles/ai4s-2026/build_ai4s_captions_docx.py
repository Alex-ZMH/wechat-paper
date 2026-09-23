from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(r"C:\Users\ironman\Desktop\wechat\wechat-article-writer\articles\ai4s-2026")
OUT = BASE / "AI4S-四张图片配文.docx"

ITEMS = [
    {
        "kicker": "配套文案 01｜调研总览",
        "title": "AI4S 调研，不只看模型有多大，更看能否进入产业链",
        "image": BASE / "images" / "ai4s-cover-industry-chain-v2.png",
        "alt": "调研 AI4S 团队与公司：从科学突破到产业落地",
        "body": "很多 AI4S 讨论停留在模型精度、论文数量和融资消息上，但产业真正关心的是：它能否把科学问题转成可执行的研发任务，并在真实工厂里稳定交付。调研 AI4S 团队与公司，建议沿着一条完整链路观察：团队是否同时理解科学机理、数据工程和生产流程；模型输出能否被实验、仿真和设备系统调用；项目是否有清晰的客户指标、验证周期和验收标准；数据与知识产权能否沉淀为下一次订单的壁垒。真正有产业价值的 AI4S，不是替科学家按下一个“生成”按钮，而是把候选配方、实验计划、工艺窗口和质量证据串起来，缩短从发现到量产的距离。本次调研将重点看团队与公司的落地位置：它们解决哪一段瓶颈，依靠什么收费，如何从一次项目走向持续复购。",
    },
    {
        "kicker": "配套文案 02｜产业结构",
        "title": "AI4S 的五层落地栈：每一层都要能被产业调用",
        "image": BASE / "images" / "ai4s-industry-stack-3d.png",
        "alt": "AI4S 产业落地五层栈",
        "body": "把 AI4S 画成五层栈，最底层是数据与算力，决定模型能否持续学习；第二层是物理化学模型，把守恒、反应机理和边界条件纳入推理；第三层是研发软件，让模型建议进入配方设计、仿真、工艺管理等日常工具；第四层是实验自动化，把计划转成可追踪、可复现的操作；最上层则是材料与药物 IP，最终要沉淀成配方、专利、工艺包或候选分子。五层之间不是简单叠加，而是相互校验：数据支撑模型，模型减少试错，软件连接组织，自动化提供证据，IP 承接商业回报。如果其中一层缺席，前面的能力就可能停留在演示或试验阶段，无法被客户稳定采用。看一家 AI4S 公司是否具备产业化能力，不能只问模型准不准，还要问它能否打通这五层，并在客户现场形成闭环。",
    },
    {
        "kicker": "配套文案 03｜商业闭环",
        "title": "模型给出方向，产业闭环决定价值如何兑现",
        "image": BASE / "images" / "ai4s-model-to-cash-loop-3d.png",
        "alt": "从模型到回款的 AI4S 产业闭环",
        "body": "AI4S 的价值，不在于模型单次给出一个漂亮答案，而在于能否把答案推进到客户愿意验收、愿意复购的结果。一个完整的产业闭环，通常从客户的科学或工程问题开始：明确指标、成本和交付边界；随后由模型检索知识、生成候选方案并给出不确定性；湿实验验证可行性，工程团队再把参数放大到设备和产线；经过试产、质量评估和客户验收，结果才真正进入订单。每一步都要有可追溯的数据、版本和责任人，失败也要沉淀为下一轮学习的样本。真正的难点，是把跨部门协作、质量体系和时间节点一起纳入流程，而不是把模型当成孤立工具。对团队的调研，关键不是看演示多流畅，而是看它如何跨过这些“闸门”，把一次性研发项目变成稳定的交付流程，并最终体现为回款与复购。",
    },
    {
        "kicker": "配套文案 04｜材料案例",
        "title": "高镍材料的产业化路径：从一组候选配方到可复购订单",
        "image": BASE / "images" / "ai4s-high-nickel-landing-3d.png",
        "alt": "高镍材料从候选到订单的产业化路径",
        "body": "以高镍正极材料为例，产业化不是从“让模型找配方”开始，而是从客户指标开始：能量密度、循环寿命、安全边界、成本和设备兼容性必须先被定义。模型据此生成一组候选配方，并说明元素比例、工艺窗口和风险假设；实验团队完成合成、粒径与晶体结构表征，筛掉不稳定方案；入选样品还要经过烧结中试，验证放大后的温度、气氛、批次一致性和良率，同时评估原料波动、设备维护与安全合规。还要把供应链稳定性和财务模型纳入评估，确认方案可以复制、交付和定价。随后进入电芯测试，观察倍率、循环、产气等关键指标，最后由客户按自己的标准验收。只有当数据可复现、工艺可转移、成本可接受，候选材料才可能从实验室资产变成订单，并通过持续改进形成下一轮复购。",
    },
]

NAVY = RGBColor(12, 38, 68)
BLUE = RGBColor(46, 116, 181)
CYAN = RGBColor(36, 167, 210)
MUTED = RGBColor(96, 108, 122)
BLACK = RGBColor(31, 34, 38)


def set_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_image_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h3._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(31, 77, 120)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)


def add_document_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("微信公众号发布素材")
    set_font(r, size=9.5, bold=True, color=CYAN)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("AI4S 产业落地图片配文")
    set_font(r, size=23, bold=True, color=NAVY)


def add_item(doc, item, index):
    if index == 0:
        add_document_title(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0 if index == 0 else 6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(item["kicker"])
    set_font(r, size=9.5, bold=True, color=CYAN)

    title = doc.add_paragraph(style="Heading 1")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run(item["title"])
    set_font(r, size=16, bold=True, color=NAVY)

    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(0)
    image_p.paragraph_format.space_after = Pt(10)
    image_p.paragraph_format.keep_with_next = True
    width = Inches(6.25)
    shape = image_p.add_run().add_picture(str(item["image"]), width=width)
    set_image_alt(shape, item["alt"])

    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(4)
    label.paragraph_format.keep_with_next = True
    r = label.add_run("配文")
    set_font(r, size=10, bold=True, color=BLUE)

    body = doc.add_paragraph()
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.10
    body.paragraph_format.keep_together = True
    r = body.add_run(item["body"])
    set_font(r, size=11, color=BLACK)


def build():
    for item in ITEMS:
        if not item["image"].exists():
            raise FileNotFoundError(item["image"])

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header_p.add_run("AI4S 产业落地调研｜图片配文")
    set_font(r, size=8.5, color=MUTED)
    header_p.paragraph_format.space_after = Pt(0)

    footer_p = section.footer.paragraphs[0]
    add_page_field(footer_p)

    for i, item in enumerate(ITEMS):
        if i:
            doc.add_page_break()
        add_item(doc, item, i)

    core = doc.core_properties
    core.title = "AI4S 产业落地图片配文"
    core.subject = "微信公众号四张图片配套文案"
    core.author = ""
    core.keywords = "AI4S, 产业落地, 微信公众号, 图片配文"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
